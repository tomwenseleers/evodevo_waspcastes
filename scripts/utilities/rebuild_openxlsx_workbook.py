from __future__ import annotations

import argparse
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter


MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"


def cell_value(cell: ET.Element, shared_strings: list[str]):
    value_node = cell.find(f"{{{MAIN}}}v")
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        return "".join(node.text or "" for node in cell.findall(f".//{{{MAIN}}}t"))
    if value_node is None or value_node.text is None:
        return None
    value = value_node.text
    if cell_type == "s":
        return shared_strings[int(value)]
    if cell_type in {"str", "e"}:
        return value
    if cell_type == "b":
        return value == "1"
    try:
        number = float(value)
        return int(number) if number.is_integer() else number
    except ValueError:
        return value


def read_sheet_rows(archive: zipfile.ZipFile, sheet_path: str, shared_strings: list[str]):
    root = ET.fromstring(archive.read(sheet_path))
    rows: list[list[object]] = []
    for row in root.findall(f".//{{{MAIN}}}sheetData/{{{MAIN}}}row"):
        values: dict[int, object] = {}
        for cell in row.findall(f"{{{MAIN}}}c"):
            reference = cell.attrib.get("r", "A1")
            match = re.match(r"([A-Z]+)", reference)
            if match is None:
                continue
            column = column_index_from_string(match.group(1))
            values[column] = cell_value(cell, shared_strings)
        if values:
            width = max(values)
            rows.append([values.get(column) for column in range(1, width + 1)])
    return rows


def rebuild(source: Path, destination: Path) -> None:
    with zipfile.ZipFile(source) as archive:
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in root.findall(f"{{{MAIN}}}si"):
                shared_strings.append(
                    "".join(node.text or "" for node in item.findall(f".//{{{MAIN}}}t"))
                )

        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {
            relation.attrib["Id"]: relation.attrib["Target"]
            for relation in relationships.findall(f"{{{PKG_REL}}}Relationship")
        }
        sheets = []
        for sheet in workbook.findall(f".//{{{MAIN}}}sheet"):
            relationship_id = sheet.attrib[f"{{{REL}}}id"]
            target = targets[relationship_id].replace("\\", "/")
            if target.startswith("/"):
                sheet_path = target.lstrip("/")
            elif target.startswith("xl/"):
                sheet_path = target
            else:
                sheet_path = f"xl/{target}"
            sheets.append((sheet.attrib["name"], sheet_path))

        output = Workbook()
        output.remove(output.active)
        header_fill = PatternFill("solid", fgColor="4F81BD")
        header_font = Font(color="FFFFFF", bold=True)
        for name, sheet_path in sheets:
            worksheet = output.create_sheet(name[:31])
            rows = read_sheet_rows(archive, sheet_path, shared_strings)
            for row in rows:
                worksheet.append(row)
            if rows:
                worksheet.freeze_panes = "A2"
                worksheet.auto_filter.ref = worksheet.dimensions
                for cell in worksheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                sample_limit = min(200, worksheet.max_row)
                for column in range(1, worksheet.max_column + 1):
                    values = [
                        str(worksheet.cell(row, column).value or "")
                        for row in range(1, sample_limit + 1)
                    ]
                    worksheet.column_dimensions[get_column_letter(column)].width = min(
                        55, max(10, max(map(len, values), default=0) + 2)
                    )

    destination.parent.mkdir(parents=True, exist_ok=True)
    output.save(destination)
    checked = load_workbook(destination, read_only=False, data_only=True)
    if len(checked.sheetnames) != len(sheets):
        raise RuntimeError("Worksheet count changed during workbook rebuild")
    if any(checked[name].max_row < 1 for name, _ in sheets):
        raise RuntimeError("A worksheet was lost during workbook rebuild")


parser = argparse.ArgumentParser()
parser.add_argument("source", type=Path)
parser.add_argument("destination", type=Path)
args = parser.parse_args()
rebuild(args.source, args.destination)
